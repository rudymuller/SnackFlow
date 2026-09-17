from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documentation" / "Manual_do_Usuario_SnackFlow.docx"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(text)
    return paragraph


def add_note(document, title, text):
    table = document.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF3F5")
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(f"{title}: ")
    run.bold = True
    paragraph.add_run(text)
    document.add_paragraph()


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in (("Title", 28, "12343B"), ("Heading 1", 17, "12343B"), ("Heading 2", 13, "237C86")):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True


def add_cover(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(65)
    run = paragraph.add_run("SNACKFLOW")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(31)
    run.font.color.rgb = RGBColor(18, 52, 59)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Manual do Usuário")
    run.font.name = "Aptos Display"
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor(35, 124, 134)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    run = paragraph.add_run("Pedidos, estoque e controle em um só lugar")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(86, 114, 134)

    table = document.add_table(rows=1, cols=1)
    table.alignment = 1
    cell = table.cell(0, 0)
    set_cell_shading(cell, "EAF3F5")
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].add_run("Versão 1.0 | Setembro de 2026")
    document.add_page_break()


def add_contents(document):
    add_heading(document, "Conteúdo", 1)
    for item in (
        "1. Sobre o SnackFlow",
        "2. Início e login",
        "3. Perfis de acesso",
        "4. Menu e navegação",
        "5. Gestão de pedidos",
        "6. Módulos administrativos",
        "7. Boas práticas e solução de problemas",
        "8. Encerramento e dados",
    ):
        document.add_paragraph(item)
    add_note(document, "Importante", "Os nomes dos botões e campos podem variar ligeiramente conforme a versão instalada. Siga as mensagens exibidas pelo sistema.")


def add_access_table(document):
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ("Perfil", "Acesso", "Uso principal")
    for index, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], text, bold=True, color="FFFFFF")
        set_cell_shading(table.rows[0].cells[index], "237C86")
    rows = (
        ("Administrador", "Todos os módulos", "Operação, cadastros, estoque e relatórios"),
        ("Funcionário", "Home e Pedidos", "Registro e acompanhamento de pedidos"),
    )
    for row in rows:
        cells = table.add_row().cells
        for index, text in enumerate(row):
            set_cell_text(cells[index], text)
    document.add_paragraph()


def build_document():
    document = Document()
    configure_document(document)
    add_cover(document)
    add_contents(document)

    add_heading(document, "1. Sobre o SnackFlow", 1)
    document.add_paragraph("O SnackFlow é um sistema desktop para apoiar a operação de lanchonetes. Ele centraliza pedidos, cadastro de lanches, estoque, contas a pagar e consultas financeiras.")
    add_bullet(document, "Use o sistema com o perfil correspondente à sua função.")
    add_bullet(document, "Mantenha os dados cadastrados atualizados para que pedidos e relatórios sejam confiáveis.")
    add_bullet(document, "O sistema utiliza um banco local SQLite; faça cópias de segurança periódicas.")

    add_heading(document, "2. Início e login", 1)
    add_heading(document, "2.1 Abrir o sistema", 2)
    add_number(document, "Execute o SnackFlow pelo atalho ou pelo arquivo do aplicativo.")
    add_number(document, "Na tela inicial, selecione Entrar.")
    add_number(document, "Informe o usuário e a senha e selecione Entrar novamente. Também é possível confirmar com a tecla Enter.")
    add_number(document, "Após a validação, o sistema abrirá o menu do seu perfil.")
    add_note(document, "Acesso inicial", "Quando não houver usuário cadastrado, os acessos internos de demonstração são `admin` com senha `admin` e `atend` com senha `atend`. Altere ou substitua esses acessos antes do uso real.")

    add_heading(document, "2.2 Cancelar ou sair do login", 2)
    add_bullet(document, "Selecione Cancelar para retornar à tela inicial.")
    add_bullet(document, "Para encerrar a sessão depois de entrar, selecione Sair da conta na barra lateral e confirme a saída.")

    add_heading(document, "3. Perfis de acesso", 1)
    add_access_table(document)
    document.add_paragraph("O administrador deve evitar compartilhar sua conta. O funcionário deve solicitar ao administrador o cadastro ou a correção de seus dados de acesso.")

    add_heading(document, "4. Menu e navegação", 1)
    document.add_paragraph("A barra lateral permanece visível durante o uso. Selecione uma opção para abrir o módulo no painel principal.")
    add_bullet(document, "Home: retorna à visão geral e mostra o usuário conectado.")
    add_bullet(document, "Pedidos: cria, consulta, atualiza e acompanha pedidos.")
    add_bullet(document, "Lanches: cadastra produtos, preços e receitas.")
    add_bullet(document, "Usuários: administra contas e perfis de acesso.")
    add_bullet(document, "Estoque: controla itens, quantidades, unidades e validade.")
    add_bullet(document, "Gastos: registra contas a pagar e seus status.")
    add_bullet(document, "Faturamento: consulta vendas por mês e detalhes de pedidos fechados.")
    add_bullet(document, "Gestão: consulta o balanço mensal de entradas, saídas e saldo.")
    add_note(document, "Seleção de registros", "Em telas com tabelas, selecione primeiro a linha desejada antes de editar, excluir ou consultar detalhes.")

    add_heading(document, "5. Gestão de pedidos", 1)
    add_heading(document, "5.1 Criar um pedido", 2)
    add_number(document, "Abra Pedidos na barra lateral.")
    add_number(document, "Selecione Novo pedido.")
    add_number(document, "Informe o nome do cliente e, se necessário, uma observação.")
    add_number(document, "Escolha a categoria e o item. Informe uma quantidade maior que zero e adicione o item ao pedido.")
    add_number(document, "Confira os itens e o valor total.")
    add_number(document, "Salve o pedido.")
    add_note(document, "Estoque", "O sistema verifica a disponibilidade dos itens e das receitas. Se faltar estoque, revise a quantidade ou peça ao administrador para atualizar o estoque.")

    add_heading(document, "5.2 Acompanhar e atualizar", 2)
    add_bullet(document, "Os pedidos podem aparecer nos estados Aberto, Em Produção, Em Consumo, Fechado ou Cancelado.")
    add_bullet(document, "Use Atualizar para alterar os dados permitidos do pedido e seu andamento.")
    add_bullet(document, "Use Excluir somente quando necessário e confirme a operação.")
    add_bullet(document, "Feche o pedido apenas após concluir o atendimento e conferir o valor.")

    add_heading(document, "6. Módulos administrativos", 1)
    add_heading(document, "6.1 Lanches", 2)
    add_number(document, "Abra Lanches e selecione Adicionar ou a opção equivalente de novo cadastro.")
    add_number(document, "Informe nome e preço de venda.")
    add_number(document, "Selecione os ingredientes, informe a quantidade e adicione cada componente à receita.")
    add_number(document, "Salve o lanche e confira a listagem.")
    add_bullet(document, "Selecione um lanche para editar ou desativar.")
    add_bullet(document, "Use Consultar estoque para verificar os ingredientes disponíveis para a receita.")

    add_heading(document, "6.2 Estoque", 2)
    add_number(document, "Abra Estoque e selecione Adicionar item.")
    add_number(document, "Preencha nome, categoria, unidade, quantidade, preço e validade quando aplicável.")
    add_number(document, "Selecione Salvar.")
    add_bullet(document, "Use Editar após selecionar um item.")
    add_bullet(document, "Use Excluir para remover ou desativar um item, confirmando a mensagem exibida.")
    add_bullet(document, "Consulte a visão agrupada para analisar quantidades consolidadas.")

    add_heading(document, "6.3 Usuários", 2)
    add_number(document, "Abra Usuários e selecione Adicionar.")
    add_number(document, "Preencha os dados solicitados e escolha o tipo de acesso adequado.")
    add_number(document, "Selecione Salvar e confirme a mensagem de sucesso.")
    add_bullet(document, "Revise cuidadosamente o perfil: Administrador possui acesso aos módulos administrativos; Funcionário possui acesso operacional aos pedidos.")

    add_heading(document, "6.4 Gastos", 2)
    add_number(document, "Abra Gastos / Contas a pagar.")
    add_number(document, "Informe descrição, categoria, valor, vencimento e status.")
    add_number(document, "Salve o registro e mantenha o status atualizado após o pagamento.")

    add_heading(document, "6.5 Faturamento", 2)
    add_number(document, "Abra Faturamento.")
    add_number(document, "Selecione o mês desejado.")
    add_number(document, "Consulte as vendas agrupadas por data.")
    add_number(document, "Selecione uma data para visualizar os pedidos fechados e, depois, os detalhes de cada pedido.")

    add_heading(document, "6.6 Gestão / balanço mensal", 2)
    add_number(document, "Abra Gestão.")
    add_number(document, "Informe o período usando os campos De e Até.")
    add_number(document, "Analise a tabela de entradas, saídas e saldo.")
    add_number(document, "Use o gráfico para identificar a evolução das movimentações.")

    add_heading(document, "7. Boas práticas e solução de problemas", 1)
    add_bullet(document, "Confira nomes, quantidades, valores e datas antes de salvar.")
    add_bullet(document, "Não encerre o computador durante uma operação de gravação.")
    add_bullet(document, "Quando um formulário exigir seleção, selecione a linha antes de clicar em editar, excluir ou consultar.")
    add_bullet(document, "Se as credenciais forem recusadas, confirme usuário e senha e solicite ao administrador a verificação do cadastro.")
    add_bullet(document, "Se um pedido informar estoque insuficiente, cadastre ou ajuste o lote correspondente em Estoque.")
    add_bullet(document, "Se um relatório estiver vazio, confira o período e se existem pedidos fechados no intervalo.")
    add_bullet(document, "Se o aplicativo não abrir, confirme que a pasta data está presente e que o banco está em data/SysDB.db.")
    add_note(document, "Mensagens de erro", "Leia a mensagem exibida, corrija o dado indicado e tente novamente. Registre o módulo, a operação e a mensagem para encaminhar um problema ao responsável técnico.")

    add_heading(document, "8. Encerramento e dados", 1)
    document.add_paragraph("Ao terminar o trabalho, selecione Sair da conta e confirme. Para distribuição do aplicativo, o banco deve ficar em `data/SysDB.db` dentro da pasta do programa.")
    add_heading(document, "8.1 Execução em ambiente de desenvolvimento", 2)
    add_bullet(document, "Com o ambiente virtual ativo, execute `python src/Main.py` a partir da pasta do projeto.")
    add_heading(document, "8.2 Versão distribuída", 2)
    add_bullet(document, "No Windows, o script de build gera o aplicativo em `dist\\SnackFlow\\SnackFlow.exe`.")
    add_bullet(document, "No Linux, o script de build gera o aplicativo em `dist/SnackFlow/SnackFlow`.")
    add_bullet(document, "Faça backup do arquivo `data/SysDB.db` antes de atualizações ou manutenções.")

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("SnackFlow | Manual do Usuário")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
